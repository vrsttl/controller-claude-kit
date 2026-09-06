# /// script
# requires-python = ">=3.12"
# dependencies = ["pandas>=2.2", "openpyxl>=3.1", "pyyaml>=6"]
# ///
"""Validation checks V01..V16, delivery and the run log.

Public API:
    run_checks(conn, frames, result, spec, period, workbook_path=None) -> list[CheckResult]
    build_run_info(spec, period, frames, result, checks, run_ts=None) -> dict
        (totals: rev_net, inv_count, ar_balance, overdue_amt)
    deliver(staging_file, spec, period, checks, run_info=None) -> DeliveryOutcome
    append_runlog(spec, period, run_info, outcome, delivered_path) -> Path
    last_runlog_entry(spec, period_label) -> dict | None
    preview_lines(spec, period, frames, result, checks, previous) -> list[str]

Tolerance semantics: a difference passes when abs(delta) < tolerance (a
tolerance of 1 HUF therefore rejects a 1 HUF delta); tolerance 0 means exact.
"""

from __future__ import annotations

import json
import shutil
import sqlite3
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

import pandas as pd
from kit_meta import KIT_VERSION
from report_engine import Frames, Period, Result, ar_frame
from report_excel import (
    MAX_EXEC_EXCEPTIONS,
    TILE_COLS,
    expected_sheets,
    expected_tables,
    sheet_name,
)
from report_spec import CHECK_IDS, Spec, spec_hash, validate

CHECK_TITLES_HU = {
    "V01": "Spec értelmezhető",
    "V02": "Forrás frissesség",
    "V03": "NAV digest darabszám = gyorsítótár",
    "V04": "Tételsorok összege = fejléc",
    "V05": "ÁFA-bontás összege = fejléc",
    "V06": "NAV HUF vs összeg × árfolyam",
    "V07": "Áfalista egyeztetés",
    "V08": "Főkönyvi CSV fizetettség egyeztetés",
    "V09": "Nincs vevő nélküli számla",
    "V10": "Lezárt időszak stabilitása",
    "V11": "Számlalánc integritás",
    "V12": "Árfolyam ésszerűség",
    "V13": "Duplikált számlák",
    "V14": "Minimális sorszám",
    "V15": "Belső egyeztetés (vevő = ÁFA-kulcs = összesen)",
    "V16": "Kimeneti fájl épsége",
}
OUTCOME_HU = {
    "delivered": "kézbesítve",
    "rejected": "elutasítva",
    "pending": "függőben (a célfájl zárolva)",
    "exists": "nem írható felül (a célfájl létezik)",
    "dry_run": "próbafuttatás",
    "not_delivered": "nem kézbesítve (--no-deliver)",
}


@dataclass
class CheckResult:
    id: str
    title_hu: str
    status: str  # ok | warn | fail | skipped
    detail_hu: str
    blocking: bool

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class DeliveryOutcome:
    exit_code: int
    outcome: str  # delivered | rejected | pending | exists
    delivered_path: Path | None
    message_hu: str


def _within(delta: float, tol: float) -> bool:
    return abs(delta) < max(float(tol), 1e-9)


def _status(spec: Spec, cid: str, failed: bool) -> str:
    if not failed:
        return "ok"
    return "fail" if spec.check(cid).severity == "fail" else "warn"


def _mk(spec: Spec, cid: str, failed: bool, detail: str, skipped: bool = False) -> CheckResult:
    status = "skipped" if skipped else _status(spec, cid, failed)
    return CheckResult(cid, CHECK_TITLES_HU[cid], status, detail, spec.check(cid).blocks_delivery)


def _period_docs(frames: Frames, period: Period) -> pd.DataFrame:
    inv = frames.invoice
    s, e = pd.Timestamp(period.start), pd.Timestamp(period.end)
    return inv[(inv["basis_date"] >= s) & (inv["basis_date"] <= e)]


def _tol_for(row, tol: float) -> float:
    return 0.01 if row.currency != "HUF" and tol >= 0.01 else tol


# ---------------------------------------------------------------------------
# Individual checks
# ---------------------------------------------------------------------------


def _v02(conn: sqlite3.Connection, spec: Spec, period: Period) -> CheckResult:
    required = [n for n in ("agent", "nav_digest") if getattr(spec.sources, n).required]
    if not required:
        return _mk(spec, "V02", False, "nincs kötelező szinkronforrás", skipped=True)
    problems = []
    for src in required:
        rows = conn.execute(
            "SELECT finished, window_from, window_to FROM sync_log WHERE source = ? AND finished "
            "IS NOT NULL "
            "AND errors = 0 ORDER BY finished DESC",
            (src,),
        ).fetchall()
        covering = [
            r
            for r in rows
            if (r[1] or "0000") <= period.start.isoformat()
            and (r[2] or "9999") >= period.end.isoformat()
        ]
        if not covering:
            problems.append(f"{src}: nincs sikeres szinkron a(z) {period.label} időszakra")
            continue
        finished = datetime.fromisoformat(covering[0][0]).date()
        age = (period.run_date - finished).days
        if age > spec.sources.max_age_days:
            problems.append(
                f"{src}: utolsó sikeres szinkron {finished.isoformat()}, {age} napos (limit "
                f"{spec.sources.max_age_days})"
            )
    return _mk(
        spec,
        "V02",
        bool(problems),
        "; ".join(problems) if problems else f"friss: {', '.join(required)}",
    )


def _v03(conn: sqlite3.Connection, frames: Frames, spec: Spec, period: Period) -> CheckResult:
    if not spec.sources.nav_digest.required:
        return _mk(spec, "V03", False, "NAV digest nem kötelező", skipped=True)
    docs = _period_docs(frames, period)
    numbers = set(docs["invoice_number"])
    rows = conn.execute(
        "SELECT DISTINCT doc_key FROM raw_documents WHERE source = 'nav_digest'"
    ).fetchall()
    nav = {r[0] for r in rows} & numbers
    missing = len(numbers) - len(nav)
    return _mk(spec, "V03", missing != 0, f"NAV digest: {len(nav)}, gyorsítótár: {len(numbers)}")


def _v04(frames: Frames, spec: Spec, period: Period) -> CheckResult:
    tol = spec.check("V04").tolerance
    docs = _period_docs(frames, period)
    lines = frames.line[frames.line["invoice_number"].isin(set(docs["invoice_number"]))]
    if lines.empty:
        return _mk(spec, "V04", False, "nincs tételsor az időszakban", skipped=True)
    sums = lines.groupby("invoice_number").agg(net=("net", "sum"), vat=("vat", "sum"))
    bad = []
    for row in docs.itertuples():
        if row.invoice_number not in sums.index:
            continue
        t = _tol_for(row, tol)
        d_net = float(sums.at[row.invoice_number, "net"]) - float(row.net_amount)
        d_vat = float(sums.at[row.invoice_number, "vat"]) - float(row.vat_amount)
        if not (_within(d_net, t) and _within(d_vat, t)):
            bad.append(f"{row.invoice_number} (nettó {d_net:+.2f}, ÁFA {d_vat:+.2f})")
    return _mk(
        spec,
        "V04",
        bool(bad),
        ("eltérés: " + ", ".join(bad[:10])) if bad else f"{len(sums)} számla egyezik",
    )


def _v05(frames: Frames, spec: Spec, period: Period) -> CheckResult:
    tol = spec.check("V05").tolerance
    docs = _period_docs(frames, period)
    vat = frames.vat[frames.vat["invoice_number"].isin(set(docs["invoice_number"]))]
    if vat.empty:
        return _mk(spec, "V05", False, "nincs ÁFA-bontás az időszakban", skipped=True)
    sums = vat.groupby("invoice_number").agg(net=("net", "sum"), vat=("vat", "sum"))
    bad = []
    for row in docs.itertuples():
        if row.invoice_number not in sums.index:
            continue
        t = _tol_for(row, tol)
        d_net = float(sums.at[row.invoice_number, "net"]) - float(row.net_amount)
        d_vat = float(sums.at[row.invoice_number, "vat"]) - float(row.vat_amount)
        if not (_within(d_net, t) and _within(d_vat, t)):
            bad.append(f"{row.invoice_number} (nettó {d_net:+.2f}, ÁFA {d_vat:+.2f})")
    return _mk(
        spec,
        "V05",
        bool(bad),
        ("eltérés: " + ", ".join(bad[:10])) if bad else f"{len(sums)} számla egyezik",
    )


def _v06(conn: sqlite3.Connection, frames: Frames, spec: Spec, period: Period) -> CheckResult:
    if not spec.sources.nav_digest.required:
        return _mk(spec, "V06", False, "NAV adat nélkül nem értelmezhető", skipped=True)
    tol = spec.check("V06").tolerance
    docs = _period_docs(frames, period)
    fx = docs[(docs["currency"] != "HUF") & (docs["net_amount"] != 0)]
    bad = []
    for row in fx.itertuples():
        dev = float(row.net_huf) / (float(row.net_amount) * float(row.fx_rate_invoice)) - 1
        if abs(dev) > tol:
            bad.append(f"{row.invoice_number} ({dev:+.2%})")
    return _mk(
        spec,
        "V06",
        bool(bad),
        ("eltérés: " + ", ".join(bad[:10])) if bad else f"{len(fx)} devizás számla rendben",
    )


def _v07(conn: sqlite3.Connection, frames: Frames, spec: Spec, period: Period) -> CheckResult:
    tol = spec.check("V07").tolerance
    st = pd.read_sql_query(
        "SELECT invoice_number, vat_rate, net, vat FROM staging_afalista WHERE issue_date BETWEEN "
        "? AND ?",
        conn,
        params=(period.start.isoformat(), period.end.isoformat()),
    )
    if st.empty:
        if spec.sources.afalista.required:
            return _mk(
                spec,
                "V07",
                True,
                f"az Áfalista kötelező, de nincs betöltött sor a(z) {period.label} időszakra",
            )
        return _mk(spec, "V07", False, "nincs Áfalista sor az időszakra", skipped=True)
    docs = _period_docs(frames, period)
    vat = frames.vat[frames.vat["invoice_number"].isin(set(docs["invoice_number"]))]
    ours = vat.groupby(["invoice_number", "vat_rate"]).agg(
        net=("net_huf", "sum"), vat=("vat_huf", "sum")
    )
    theirs = (
        st.assign(vat_rate=st["vat_rate"].astype(str))
        .groupby(["invoice_number", "vat_rate"])
        .agg(net=("net", "sum"), vat=("vat", "sum"))
    )
    bad = []
    for key, row in theirs.iterrows():
        if key not in ours.index:
            bad.append(f"{key[0]}/{key[1]} hiányzik a gyorsítótárból")
            continue
        d_net = float(row["net"]) - float(ours.at[key, "net"])
        d_vat = float(row["vat"]) - float(ours.at[key, "vat"])
        if not (_within(d_net, tol) and _within(d_vat, tol)):
            bad.append(f"{key[0]}/{key[1]} (nettó {d_net:+.0f}, ÁFA {d_vat:+.0f})")
    return _mk(
        spec,
        "V07",
        bool(bad),
        ("eltérés: " + ", ".join(bad[:10])) if bad else f"{len(theirs)} sor egyezik",
    )


_STATUS_MAP = {
    "fizetve": "paid",
    "kifizetve": "paid",
    "paid": "paid",
    "nem fizetve": "unpaid",
    "kifizetetlen": "unpaid",
    "unpaid": "unpaid",
    "részben fizetve": "partial",
    "reszben fizetve": "partial",
    "partial": "partial",
}


def _v08(conn: sqlite3.Connection, frames: Frames, spec: Spec, period: Period) -> CheckResult:
    st = pd.read_sql_query(
        "SELECT DISTINCT invoice_number, pay_status_raw FROM staging_fokonyvi WHERE issue_date "
        "BETWEEN ? AND ?",
        conn,
        params=(period.start.isoformat(), period.end.isoformat()),
    )
    if st.empty:
        if spec.sources.fokonyvi_csv.required:
            return _mk(
                spec,
                "V08",
                True,
                f"a főkönyvi CSV kötelező, de nincs betöltött sor a(z) {period.label} időszakra",
            )
        return _mk(spec, "V08", False, "nincs főkönyvi CSV sor az időszakra", skipped=True)
    derived = frames.invoice.set_index("invoice_number")["pay_status_raw"]
    bad = []
    checked = 0
    for row in st.itertuples():
        raw = str(row.pay_status_raw or "").strip().lower()
        mapped = _STATUS_MAP.get(raw)
        if mapped is None or row.invoice_number not in derived.index:
            continue
        checked += 1
        if derived[row.invoice_number] != mapped:
            bad.append(
                f"{row.invoice_number} (CSV: {mapped}, gyorsítótár: {derived[row.invoice_number]})"
            )
    return _mk(
        spec,
        "V08",
        bool(bad),
        ("eltérés: " + ", ".join(bad[:10])) if bad else f"{checked} számla egyezik",
    )


def _v09(frames: Frames, spec: Spec) -> CheckResult:
    inv = frames.invoice
    names = frames.customer.set_index("customer_key")["name"]
    missing = inv[
        ~inv["customer_key"].isin(names.index)
        | (inv["customer_key"].map(names).fillna("").astype(str).str.strip() == "")
    ]
    return _mk(
        spec,
        "V09",
        len(missing) > 0,
        ("vevő nélkül: " + ", ".join(missing["invoice_number"].head(10)))
        if len(missing)
        else "minden számlához van vevő",
    )


def _v10(spec: Spec, period: Period, result: Result, frames: Frames) -> CheckResult:
    prev = last_runlog_entry(spec, period.label)
    if prev is None:
        return _mk(spec, "V10", False, "első futtatás erre az időszakra")
    cur = {"rev_net": _rev_net_value(frames, period), "inv_count": _inv_count_value(frames, period)}
    prev_vals = prev.get("totals") or {}
    diffs = []
    for k, v in cur.items():
        pv = prev_vals.get(k)
        if pv is not None and not _within(float(v) - float(pv), 1.0):
            diffs.append(f"{k}: {pv} -> {v}")
    return _mk(
        spec,
        "V10",
        bool(diffs),
        ("változás az előző futtatáshoz képest: " + "; ".join(diffs))
        if diffs
        else "egyezik az előző futtatással",
    )


def _v11(conn: sqlite3.Connection, spec: Spec) -> CheckResult:
    rows = conn.execute(
        "SELECT i.invoice_number FROM invoice i LEFT JOIN invoice r ON r.invoice_number = "
        "i.chain_root "
        "WHERE i.doc_type IN ('modifier','storno') AND r.invoice_number IS NULL"
    ).fetchall()
    bad = [r[0] for r in rows]
    return _mk(
        spec,
        "V11",
        bool(bad),
        ("hiányzó eredeti számla: " + ", ".join(bad[:10]))
        if bad
        else "minden helyesbítő és sztornó lánca teljes",
    )


def _v12(frames: Frames, spec: Spec, period: Period) -> CheckResult:
    factor = spec.check("V12").tolerance or 2.0
    docs = _period_docs(frames, period)
    fx = docs[docs["currency"] != "HUF"]
    if fx.empty:
        return _mk(spec, "V12", False, "nincs devizás számla az időszakban", skipped=True)
    # reference median comes from the whole loaded history, so a two-invoice
    # period cannot skew its own median
    hist = frames.invoice[frames.invoice["currency"] != "HUF"]
    bad = []
    for cur, grp in fx.groupby("currency"):
        med = float(hist.loc[hist["currency"] == cur, "fx_rate_invoice"].median())
        for row in grp.itertuples():
            r = float(row.fx_rate_invoice)
            if med > 0 and not (med / factor <= r <= med * factor):
                bad.append(f"{row.invoice_number} ({cur} {r} vs medián {med})")
    return _mk(
        spec,
        "V12",
        bool(bad),
        ("gyanús árfolyam: " + ", ".join(bad[:10]))
        if bad
        else f"{len(fx)} devizás számla ésszerű árfolyammal",
    )


def _v13(frames: Frames, spec: Spec, period: Period) -> CheckResult:
    docs = _period_docs(frames, period)
    inv = docs[docs["doc_type"] == "invoice"]
    dup = inv[inv.duplicated(["customer_key", "gross_huf", "issue_date"], keep=False)]
    groups = dup.groupby(["customer_key", "gross_huf", "issue_date"])["invoice_number"].apply(list)
    bad = ["/".join(v) for v in groups]
    return _mk(
        spec,
        "V13",
        bool(bad),
        ("lehetséges duplikátum: " + ", ".join(bad[:10])) if bad else "nincs duplikátum",
    )


def _v14(frames: Frames, spec: Spec, period: Period) -> CheckResult:
    n = int(len(_period_docs(frames, period)))
    return _mk(
        spec,
        "V14",
        n < spec.validation.min_rows,
        f"{n} bizonylat az időszakban (minimum {spec.validation.min_rows})",
    )


def _rev_net_value(frames: Frames, period: Period) -> float:
    docs = _period_docs(frames, period)
    return float(docs["net_huf"].sum())


def _inv_count_value(frames: Frames, period: Period) -> int:
    docs = _period_docs(frames, period)
    return int((docs["doc_type"] == "invoice").sum())


def _v15(frames: Frames, spec: Spec, period: Period) -> CheckResult:
    tol = spec.check("V15").tolerance
    docs = _period_docs(frames, period)
    total = float(docs["net_huf"].sum())
    by_cust = float(docs.groupby("customer_key")["net_huf"].sum().sum())
    vat = frames.vat[frames.vat["invoice_number"].isin(set(docs["invoice_number"]))]
    by_rate = float(vat["net_huf"].sum())
    if spec.filters.vat_rates or vat.empty:
        ok = _within(total - by_cust, tol)
        detail = (
            f"összesen {total:,.0f} = vevőnként {by_cust:,.0f} "
            "(ÁFA-bontás kihagyva: szűrt kulcsok vagy hiányzó bontás)"
        )
    else:
        ok = _within(total - by_cust, tol) and _within(total - by_rate, tol)
        detail = f"összesen {total:,.0f}, vevőnként {by_cust:,.0f}, ÁFA-kulcsonként {by_rate:,.0f}"
    return _mk(spec, "V15", not ok, detail)


def check_output(spec: Spec, workbook_path: Path | str) -> CheckResult:
    """V16: workbook opens, expected sheets and tables exist, tile and exception limits hold."""
    try:
        import warnings

        from openpyxl import load_workbook

        with warnings.catch_warnings():
            warnings.simplefilter(
                "ignore"
            )  # openpyxl drops sparkline/CF extensions it cannot parse
            wb = load_workbook(str(workbook_path), read_only=False)
    except Exception as exc:  # noqa: BLE001 - any failure to open is a fail
        return _mk(spec, "V16", True, f"a munkafüzet nem nyitható meg: {exc}")
    problems = []
    names = set(wb.sheetnames)
    for s in expected_sheets(spec):
        if s not in names:
            problems.append(f"hiányzó lap: {s}")
    tables = {t for ws in wb.worksheets for t in ws.tables}
    for t in expected_tables(spec):
        if t not in tables:
            problems.append(f"hiányzó tábla: {t}")
    exec_name = sheet_name("executive", spec.output.language)
    if exec_name in names:
        ws = wb[exec_name]
        tile_labels = [ws.cell(row=4, column=c + 1).value for c in TILE_COLS]
        tiles = sum(1 for v in tile_labels if v)
        if tiles > 6:
            problems.append(f"{tiles} csempe (max 6)")
        merged = [str(r) for r in ws.merged_cells.ranges if r.min_row != 1]
        if merged:
            problems.append("összevont cellák a címsoron kívül: " + ", ".join(merged[:5]))
        if _exec_exception_rows(ws) > MAX_EXEC_EXCEPTIONS:
            problems.append("több mint 12 kivétel a vezetői lapon")
    wb.close()
    return _mk(
        spec,
        "V16",
        bool(problems),
        "; ".join(problems) if problems else "lapok, táblák és korlátok rendben",
    )


def _exec_exception_rows(ws) -> int:
    """Count exception rows under the block header in column M (13)."""
    count = 0
    started = False
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=13).value
        if isinstance(v, str) and v.startswith(("Kivételek", "Exceptions")):
            started = True
            continue
        if started:
            if v in ("Típus", "Type"):
                continue
            if v is None or v == "":
                break
            count += 1
    return count


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------


def run_checks(
    conn: sqlite3.Connection,
    frames: Frames,
    result: Result,
    spec: Spec,
    period: Period,
    workbook_path: Path | str | None = None,
) -> list[CheckResult]:
    errors = validate(spec)
    checks: list[CheckResult] = [
        _mk(
            spec,
            "V01",
            bool(errors),
            "; ".join(f"{e.path}: {e.message_hu}" for e in errors[:5])
            if errors
            else "a spec érvényes",
        ),
        _v02(conn, spec, period),
        _v03(conn, frames, spec, period),
        _v04(frames, spec, period),
        _v05(frames, spec, period),
        _v06(conn, frames, spec, period),
        _v07(conn, frames, spec, period),
        _v08(conn, frames, spec, period),
        _v09(frames, spec),
        _v10(spec, period, result, frames),
        _v11(conn, spec),
        _v12(frames, spec, period),
        _v13(frames, spec, period),
        _v14(frames, spec, period),
        _v15(frames, spec, period),
    ]
    if workbook_path is None:
        checks.append(_mk(spec, "V16", False, "a munkafüzet még nem készült el", skipped=True))
    else:
        checks.append(check_output(spec, workbook_path))
    assert [c.id for c in checks] == list(CHECK_IDS)
    return checks


def blocking_failures(checks: list[CheckResult]) -> list[CheckResult]:
    return [c for c in checks if c.status == "fail" and c.blocking]


def _ar_totals(frames: Frames, result: Result, period: Period) -> tuple[float, float]:
    """(ar_balance, overdue_amt) as of the period's as_of date.

    Uses the computed measures when the spec has them, else the same ar_frame
    arithmetic as m_ar_balance / m_overdue_amt, so the totals exist for every spec.
    """
    ar_balance = result.value("ar_balance")
    overdue = result.value("overdue_amt")
    if ar_balance is None or overdue is None:
        ar = ar_frame(frames, period.as_of)
        if ar_balance is None:
            ar_balance = float(ar["open_huf"].sum())
        if overdue is None:
            overdue = float(
                ar.loc[(ar["days_overdue"] > 0) & (ar["open_huf"] > 0), "open_huf"].sum()
            )
    return round(float(ar_balance), 2), round(float(overdue), 2)


def build_run_info(
    spec: Spec,
    period: Period,
    frames: Frames,
    result: Result,
    checks: list[CheckResult],
    run_ts: str | None = None,
    delivered_path: Path | str | None = None,
) -> dict:
    counts = frames.row_counts()
    counts["measure"] = int(len(result.measures))
    counts["exceptions"] = int(len(result.exceptions))
    ar_balance, overdue_amt = _ar_totals(frames, result, period)
    return {
        "period": period.label,
        "run_ts": run_ts or datetime.now().strftime("%Y-%m-%d %H:%M"),
        "spec_version": spec.report.version,
        "spec_hash": spec_hash(spec.path) if spec.path else "",
        "kit_version": KIT_VERSION,
        "row_counts": counts,
        "checks": checks,
        "totals": {
            "rev_net": round(_rev_net_value(frames, period), 2),
            "inv_count": _inv_count_value(frames, period),
            "ar_balance": ar_balance,
            "overdue_amt": overdue_amt,
        },
        "delivered_path": str(delivered_path) if delivered_path else None,
    }


def output_filename(spec: Spec, period: Period) -> str:
    return spec.output.file_pattern.format(
        slug=spec.report.slug,
        period=period.label,
        version=spec.report.version,
        run_date=period.run_date.isoformat(),
    )


def runlog_path(spec: Spec) -> Path:
    return spec.spec_dir / "runlog.jsonl"


def append_runlog(
    spec: Spec, period: Period, run_info: dict, outcome: str, delivered_path: Path | str | None
) -> Path:
    path = runlog_path(spec)
    entry = {
        "period": period.label,
        "run_ts": run_info.get("run_ts"),
        "spec_version": run_info.get("spec_version"),
        "spec_hash": run_info.get("spec_hash"),
        "kit_version": run_info.get("kit_version"),
        "row_counts": run_info.get("row_counts"),
        "totals": run_info.get("totals"),
        "checks": [
            c.to_dict() if isinstance(c, CheckResult) else c for c in run_info.get("checks") or []
        ],
        "outcome": outcome,
        "delivered_path": str(delivered_path) if delivered_path else None,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(entry, ensure_ascii=False) + "\n")
    return path


def read_runlog(path: Path) -> list[dict]:
    if not path.exists():
        return []
    entries = []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            entries.append(json.loads(line))
        except json.JSONDecodeError:
            continue
    return entries


def last_runlog_entry(spec: Spec, period_label: str | None = None) -> dict | None:
    entries = read_runlog(runlog_path(spec))
    if period_label is not None:
        entries = [e for e in entries if e.get("period") == period_label]
    return entries[-1] if entries else None


def _version_suffix(spec: Spec) -> str:
    return f"_v{spec.report.version}"


def _prune(folder: Path, spec: Spec, keep: int, protect: Path) -> list[Path]:
    pattern = spec.output.file_pattern
    glob = pattern.format(slug=spec.report.slug, period="*", version="*", run_date="*")
    candidates = [
        p
        for p in folder.glob(glob)
        if p.is_file() and p.resolve() != protect.resolve() and not p.name.endswith(".pending.xlsx")
    ]
    candidates.sort(key=lambda p: (p.stat().st_mtime, p.name))
    removed = []
    surplus = len(candidates) + 1 - keep
    for p in candidates[: max(surplus, 0)]:
        p.unlink()
        removed.append(p)
    return removed


def deliver(
    staging_file: Path | str,
    spec: Spec,
    period: Period,
    checks: list[CheckResult],
    run_info: dict | None = None,
) -> DeliveryOutcome:
    staging = Path(staging_file)
    folder = spec.resolve_path(spec.delivery.folder)
    name = output_filename(spec, period)
    run_info = run_info or {}

    def log(outcome: str, path: Path | None) -> None:
        append_runlog(spec, period, run_info, outcome, path)

    failures = blocking_failures(checks)
    if failures:
        rejected_dir = folder / "_rejected"
        rejected_dir.mkdir(parents=True, exist_ok=True)
        target = rejected_dir / f"{Path(name).stem}_FAILED.xlsx"
        shutil.move(str(staging), str(target))
        ids = ", ".join(f"{c.id} ({c.detail_hu})" for c in failures)
        log("rejected", target)
        return DeliveryOutcome(
            1,
            "rejected",
            target,
            f"ELUTASÍTVA: blokkoló ellenőrzés hibázott: {ids}. A fájl ide került: {target}",
        )

    folder.mkdir(parents=True, exist_ok=True)
    target = folder / name
    mode = spec.delivery.overwrite
    if target.exists():
        if mode == "never":
            log("exists", None)
            return DeliveryOutcome(
                1,
                "exists",
                None,
                f"A célfájl már létezik és a felülírás tiltott (overwrite: never): {target}",
            )
        if mode == "same_version" and _version_suffix(spec) not in target.name:
            target = folder / f"{Path(name).stem}{_version_suffix(spec)}.xlsx"
    try:
        shutil.move(str(staging), str(target))
    except PermissionError:
        pending = folder / f"{Path(name).stem}.pending.xlsx"
        shutil.copyfile(str(staging), str(pending))
        log("pending", pending)
        return DeliveryOutcome(
            1,
            "pending",
            pending,
            "A célfájl zárolva (nyitva Excelben vagy OneDrive szinkron alatt). A riport ide "
            f"került: {pending}. Zárd be a fájlt, majd nevezd át.",
        )
    removed = _prune(folder, spec, spec.delivery.keep_n_versions, target)
    log("delivered", target)
    msg = f"Kézbesítve: {target}"
    if removed:
        msg += f" (törölt régi verziók: {', '.join(p.name for p in removed)})"
    return DeliveryOutcome(0, "delivered", target, msg)


# ---------------------------------------------------------------------------
# Preview
# ---------------------------------------------------------------------------


def _fmt_huf(v: float | None) -> str:
    return "n.a." if v is None else f"{v:,.0f} Ft".replace(",", " ")


def preview_lines(
    spec: Spec,
    period: Period,
    frames: Frames,
    result: Result,
    checks: list[CheckResult],
    previous: dict | None,
) -> list[str]:
    counts = frames.row_counts()
    docs = _period_docs(frames, period)
    rev = _rev_net_value(frames, period)
    ar, overdue = _ar_totals(frames, result, period)
    lines = [
        f"{spec.report.title_hu} · {period.label} ({period.start} .. {period.end}) · alap: "
        f"{spec.period.basis}",
        f"Sorok: {len(docs)} bizonylat az időszakban, összesen {counts['invoice']} számla, "
        f"{counts['line']} tétel, {counts['payment']} kifizetés, {counts['customer']} vevő",
        f"Nettó árbevétel: {_fmt_huf(rev)} · Kintlévőség ({period.as_of}): {_fmt_huf(ar)}, "
        f"ebből lejárt: {_fmt_huf(overdue)}",
    ]
    if previous:
        pt = previous.get("totals") or {}
        prev_rev = pt.get("rev_net")
        if prev_rev is not None:
            lines.append(
                f"Előző futtatás ({previous.get('run_ts')}): nettó {_fmt_huf(float(prev_rev))}, "
                f"eltérés {_fmt_huf(rev - float(prev_rev))}"
            )
    else:
        lines.append("Előző futtatás: nincs ehhez az időszakhoz")
    lines.append("Ellenőrzések:")
    st = {"ok": "OK  ", "warn": "FIGY", "fail": "HIBA", "skipped": "----"}
    for c in checks:
        lines.append(f"  {c.id} [{st[c.status]}] {c.title_hu}: {c.detail_hu}")
    fails = blocking_failures(checks)
    warns = [c for c in checks if c.status == "warn"]
    lines.append(f"Összesen: {len(fails)} blokkoló hiba, {len(warns)} figyelmeztetés")
    return lines


__all__ = [
    "CheckResult",
    "DeliveryOutcome",
    "run_checks",
    "check_output",
    "blocking_failures",
    "build_run_info",
    "deliver",
    "append_runlog",
    "read_runlog",
    "last_runlog_entry",
    "output_filename",
    "runlog_path",
    "preview_lines",
    "OUTCOME_HU",
]
