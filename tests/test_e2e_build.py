"""End-to-end build on the synthetic fixture (PLAN, Verification item 2).

The instantiated build.py (copied from scripts/build_template.py next to the
havi_arbev_kintlev example spec) runs as a subprocess with --db pointing at
the engine fixture database and the delivery folder under a temporary path.
"""

from __future__ import annotations

import hashlib
import json
import sqlite3
import subprocess
import sys
import warnings
from pathlib import Path

import pytest
import report_excel as rx
import report_spec as rs
import yaml
from engine_fixture import PLANTED, make_engine_db, make_project
from openpyxl import load_workbook
from report_validate import _exec_exception_rows
from test_report_excel import GOLDEN_SHA256

SLUG = "havi_arbev_kintlev"
PERIOD = "2026-08"
RUN_DATE = "2026-09-03"
DELIVERED_NAME = f"{SLUG}_{PERIOD}_v1.0.0.xlsx"
JSON_KEYS = {
    "slug",
    "period",
    "version",
    "outcome",
    "outcome_hu",
    "delivered_path",
    "staging_path",
    "checks",
    "row_counts",
    "totals",
    "previous_totals",
    "powerbi_files",
    "exit_code",
}
TOTAL_KEYS = {"rev_net", "inv_count", "ar_balance", "overdue_amt"}
CHECK_KEYS = {"id", "title_hu", "status", "detail_hu", "blocking"}


@pytest.fixture(scope="module")
def project(tmp_path_factory) -> dict:
    tmp = tmp_path_factory.mktemp("e2e")
    root = make_project(tmp / "Riportok", [SLUG])
    spec_path = root / "reports" / SLUG / "spec.yaml"
    raw = yaml.safe_load(spec_path.read_text(encoding="utf-8"))
    raw["delivery"]["folder"] = str(tmp / "delivery")
    raw["powerbi"] = {
        "enabled": True,
        "folder": str(tmp / "powerbi"),
        "files": list(rs.POWERBI_FILES),
    }
    spec_path.write_text(rs.dump_yaml(raw), encoding="utf-8")
    assert rs.validate(rs.load_spec(spec_path)) == []
    fixture_db = tmp / "fixture.db"
    make_engine_db(fixture_db).close()
    return {
        "tmp": tmp,
        "build": root / "reports" / SLUG / "build.py",
        "spec_dir": root / "reports" / SLUG,
        "db": fixture_db,
        "delivery": tmp / "delivery",
        "powerbi": tmp / "powerbi",
        "staging": tmp / "staging",
    }


def run_build(project: dict, *extra: str, out: bool = True) -> tuple[int, list[str], dict]:
    cmd = [
        sys.executable,
        str(project["build"]),
        "--period",
        PERIOD,
        "--run-date",
        RUN_DATE,
        "--db",
        str(project["db"]),
        "--json",
        *extra,
    ]
    if out:
        cmd += ["--out", str(project["staging"])]
    proc = subprocess.run(cmd, capture_output=True, text=True, cwd=str(project["spec_dir"]))
    lines = proc.stdout.splitlines()
    assert proc.returncode in (0, 1, 2), proc.stderr
    summary = json.loads(lines[-1]) if lines and lines[-1].startswith("{") else {}
    return proc.returncode, lines, summary


def _load(path: Path):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(str(path))


def runlog_entries(project: dict) -> list[dict]:
    log = project["spec_dir"] / "runlog.jsonl"
    return [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines()]


def test_first_run_delivers_and_matches_contract(project):
    rc, lines, summary = run_build(project)
    assert rc == 0 and summary["exit_code"] == 0
    assert set(summary) == JSON_KEYS
    assert (summary["slug"], summary["period"], summary["version"]) == (SLUG, PERIOD, "1.0.0")
    assert summary["outcome"] == "delivered" and summary["outcome_hu"] == "kézbesítve"
    delivered = Path(summary["delivered_path"])
    assert delivered == project["delivery"] / DELIVERED_NAME and delivered.exists()
    assert not Path(summary["staging_path"]).exists()
    assert Path(summary["staging_path"]).parent == project["staging"]

    checks = summary["checks"]
    assert [c["id"] for c in checks] == list(rs.CHECK_IDS)
    assert all(set(c) == CHECK_KEYS for c in checks)
    assert not [c for c in checks if c["status"] == "fail" and c["blocking"]]
    assert {c["id"]: c["status"] for c in checks}["V16"] == "ok"

    totals = summary["totals"]
    assert set(totals) == TOTAL_KEYS
    assert totals["rev_net"] > 0 and totals["inv_count"] == 28
    assert 0 < totals["overdue_amt"] <= totals["ar_balance"]
    assert summary["previous_totals"] is None
    assert summary["row_counts"]["invoice"] > 0 and summary["row_counts"]["exceptions"] >= 0

    # Hungarian preview lines precede the JSON line
    assert lines[0].startswith("Havi árbevétel és kintlévőség · 2026-08")
    assert any("Kintlévőség (2026-08-31)" in line and "ebből lejárt" in line for line in lines)
    assert "Előző futtatás: nincs ehhez az időszakhoz" in lines
    assert any(line.startswith("Kézbesítve: ") for line in lines)

    # Power BI CSVs: golden hashes shared with test_report_excel
    files = {Path(p).stem: Path(p) for p in summary["powerbi_files"]}
    assert set(files) == set(rs.POWERBI_FILES)
    assert all(p.parent == project["powerbi"] for p in files.values())
    hashes = {name: hashlib.sha256(p.read_bytes()).hexdigest() for name, p in files.items()}
    assert hashes == GOLDEN_SHA256

    entries = runlog_entries(project)
    assert len(entries) == 1 and entries[0]["outcome"] == "delivered"
    assert entries[0]["delivered_path"] == str(delivered) and entries[0]["period"] == PERIOD
    assert set(entries[0]["totals"]) == TOTAL_KEYS


def test_delivered_workbook_structure(project):
    delivered = project["delivery"] / DELIVERED_NAME
    wb = _load(delivered)
    spec = rs.load_spec(project["spec_dir"] / "spec.yaml")
    assert wb.sheetnames == rx.expected_sheets(spec) + [rx.SERIES_SHEET]
    assert wb.sheetnames[:2] == ["Vezetői összefoglaló", "Kivételek"]
    tables = {name for ws in wb.worksheets for name in ws.tables}
    assert tables == set(rx.expected_tables(spec))
    assert tables == {f"tbl_{SLUG}_{e}" for e in ("exceptions", "invoice", "customer", "measure")}

    ws = wb["Vezetői összefoglaló"]
    tile_labels = [ws.cell(row=4, column=c + 1).value for c in rx.TILE_COLS]
    assert 0 < sum(1 for t in tile_labels if t) <= rs.MAX_TILES
    assert ws["B5"].number_format == '#,##0," e Ft"'  # rev_net, huf_k
    assert ws["E5"].number_format == '#,##0," e Ft"'  # rev_gross
    assert ws["H5"].number_format == '#,##0," e Ft"'  # ar_balance
    assert ws["L5"].number_format == '#,##0," e Ft"'  # overdue_amt
    assert ws["O5"].number_format == '0" nap"'  # dso
    assert ws["R5"].number_format == '#,##0" db"'  # inv_count
    assert _exec_exception_rows(ws) <= rx.MAX_EXEC_EXCEPTIONS
    assert ws["P1"].value in ("ELLENŐRZÖTT", "FIGYELMEZTETÉS (1)")


def test_second_run_reports_previous_totals(project):
    rc, lines, summary = run_build(project)
    assert rc == 0 and summary["outcome"] == "delivered"
    first = runlog_entries(project)[0]["totals"]
    assert summary["previous_totals"] == first and summary["totals"] == first
    assert any(line.startswith("Előző futtatás (") and "eltérés 0 Ft" in line for line in lines)
    assert len(runlog_entries(project)) == 2
    assert (project["delivery"] / DELIVERED_NAME).exists()


def test_one_huf_vat_delta_is_rejected(project):
    doc = PLANTED["afalista_delta"].number
    conn = sqlite3.connect(project["db"])
    try:
        conn.execute(
            "UPDATE invoice_vat SET vat = vat + 1 WHERE invoice_number = ? AND vat_rate = '27'",
            (doc,),
        )
        conn.commit()
        rc, lines, summary = run_build(project)
    finally:
        conn.execute(
            "UPDATE invoice_vat SET vat = vat - 1 WHERE invoice_number = ? AND vat_rate = '27'",
            (doc,),
        )
        conn.commit()
        conn.close()
    assert rc == 1 and summary["exit_code"] == 1
    assert summary["outcome"] == "rejected" and summary["outcome_hu"] == "elutasítva"
    rejected = project["delivery"] / "_rejected" / f"{SLUG}_{PERIOD}_v1.0.0_FAILED.xlsx"
    assert Path(summary["delivered_path"]) == rejected and rejected.exists()
    v05 = next(c for c in summary["checks"] if c["id"] == "V05")
    assert v05["status"] == "fail" and v05["blocking"] and doc in v05["detail_hu"]
    assert any(line.startswith("ELUTASÍTVA") and "V05" in line for line in lines)
    assert any("1 blokkoló hiba" in line for line in lines)
    entries = runlog_entries(project)
    assert len(entries) == 3 and entries[-1]["outcome"] == "rejected"
    assert (project["delivery"] / DELIVERED_NAME).exists()  # earlier delivery untouched


def test_dry_run_writes_under_dryrun_and_never_delivers(project):
    rc, lines, summary = run_build(project, "--dry-run", out=False)
    assert rc == 0 and summary["outcome"] == "dry_run" and summary["delivered_path"] is None
    staging = Path(summary["staging_path"])
    assert staging.parent == project["spec_dir"] / "_dryrun" and staging.exists()
    assert all(
        Path(p).parent == project["spec_dir"] / "_dryrun" / "powerbi"
        for p in summary["powerbi_files"]
    )
    assert any(line.startswith("Próbafuttatás: ") for line in lines)
    assert len(runlog_entries(project)) == 3  # dry runs are not logged


def test_no_deliver_keeps_staging_file(project):
    rc, lines, summary = run_build(project, "--no-deliver", "--no-powerbi")
    assert rc == 0 and summary["outcome"] == "not_delivered"
    assert Path(summary["staging_path"]).exists() and summary["powerbi_files"] == []
    assert any(line.startswith("Nem kézbesítve") for line in lines)


def test_usage_and_spec_errors_exit_2(project):
    rc, lines, summary = run_build(project, "--period", "nonsense")
    assert rc == 2 and summary == {} and lines[-1].startswith("IDŐSZAK HIBA")
    missing = [
        sys.executable,
        str(project["build"]),
        "--period",
        PERIOD,
        "--db",
        str(project["tmp"] / "nincs.db"),
    ]
    proc = subprocess.run(missing, capture_output=True, text=True, cwd=str(project["spec_dir"]))
    assert proc.returncode == 2 and "ADATBÁZIS HIBA" in proc.stdout
