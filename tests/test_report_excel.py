"""Tests for report_excel: workbook structure read back with openpyxl, Power BI CSVs."""

from __future__ import annotations

import copy
import hashlib
import warnings
import zipfile
from datetime import date

import pytest
import report_engine as re_
import report_excel as rx
import report_spec as rs
import yaml
from engine_fixture import EXAMPLES_DIR, RUN_DATE, make_engine_db
from kit_meta import KIT_VERSION
from openpyxl import load_workbook
from report_validate import build_run_info, check_output, run_checks

# Golden sha256 of the Power BI CSVs produced from the fixture with the
# havi_arbev_kintlev example (all six files enabled). Regenerate deliberately
# when the fixture or the export columns change:
#   uv run python -c "import test_report_excel as t; t.print_golden()"
GOLDEN_SHA256 = {
    "fact_invoice": "d68c5d06bd769e5084b0ee2c8ae973fdf485cc0e057c9e39523d36d62a4e1b50",
    "fact_invoice_line": "80e1c30693dfb2b3de96f98150fd47505f787c25de914737c739b965cba9ce62",
    "fact_payment": "74a9501cd50ffd908ee1047749ba515351ffbbc21a27d0b35f5d4cbd8d16a135",
    "fact_measure": "83512ef8357948212788e113f4a39473b919691a91a5a09f90a0f0375b069c1d",
    "dim_customer": "77baea522d61e2f1287110fc26b55436c6f675c4a80dd37dbf7e00a10a96b20e",
    "dim_date": "295bdc167f7d47b73de45939300640381fe1f78badfa0cbb29e167ef27250c34",
}


def _load(path):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(str(path))


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    root = tmp_path_factory.mktemp("rip")
    spec_dir = root / "reports" / "havi_arbev_kintlev"
    spec_dir.mkdir(parents=True)
    raw = yaml.safe_load(
        (EXAMPLES_DIR / "havi_arbev_kintlev" / "spec.yaml").read_text(encoding="utf-8")
    )
    raw["powerbi"]["files"] = list(rs.POWERBI_FILES)
    spec_path = spec_dir / "spec.yaml"
    spec_path.write_text(rs.dump_yaml(raw), encoding="utf-8")
    spec = rs.load_spec(spec_path)
    assert rs.validate(spec) == []
    conn = make_engine_db(":memory:")
    period = re_.resolve_period(spec, run_date=RUN_DATE)
    frames = re_.load_frames(conn, spec, period)
    result = re_.compute(frames, spec, period)
    checks = run_checks(conn, frames, result, spec, period)
    run_info = build_run_info(spec, period, frames, result, checks, run_ts="2026-09-03 08:12")
    out = spec_dir / ".staging" / "havi.xlsx"
    rx.write_workbook(result, spec, period, out, run_info, frames=frames)
    return dict(
        spec=spec,
        period=period,
        frames=frames,
        result=result,
        checks=checks,
        run_info=run_info,
        out=out,
        conn=conn,
        root=root,
    )


def test_sheet_names_and_hidden_series(built):
    wb = _load(built["out"])
    expected = [
        "Vezetői összefoglaló",
        "Kivételek",
        "Számlák",
        "Ügyfelek",
        "Mutatók",
        "Definíciók",
        "Futtatási napló",
    ]
    assert wb.sheetnames == expected + [rx.SERIES_SHEET]
    assert wb[rx.SERIES_SHEET].sheet_state == "hidden"
    assert rx.expected_sheets(built["spec"]) == expected


def test_tables_names_and_ranges(built):
    wb = _load(built["out"])
    frames, result, spec = built["frames"], built["result"], built["spec"]
    expect_rows = {
        "exceptions": len(result.exceptions),
        "invoice": len(frames.invoice),
        "customer": len(frames.customer),
        "measure": len(rx.entity_frame("measure", result, frames, spec)),
    }
    found = {}
    for ws in wb.worksheets:
        for name, ref in ws.tables.items():
            found[name] = (ws.title, ref)
    assert (
        set(found)
        == set(rx.expected_tables(spec))
        == {f"tbl_havi_arbev_kintlev_{e}" for e in expect_rows}
    )
    for entity, rows in expect_rows.items():
        title, ref = found[f"tbl_havi_arbev_kintlev_{entity}"]
        assert title == rx.sheet_name(entity, "hu")
        assert ref.startswith("A1:") and ref.endswith(str(rows + 1))
        headers = [c.value for c in wb[title][1]]
        assert all(h == h.lower() and " " not in h for h in headers)
    assert [c.value for c in wb["Kivételek"][1]] == re_.EXCEPTION_COLUMNS


def test_executive_layout_limits_and_formats(built):
    wb = _load(built["out"])
    ws = wb["Vezetői összefoglaló"]
    tiles = [ws.cell(row=4, column=c + 1).value for c in rx.TILE_COLS]
    assert sum(1 for t in tiles if t) == 6 <= rs.MAX_TILES
    assert tiles[0] == "Nettó árbevétel" and tiles[4].startswith("DSO")
    assert ws["B5"].number_format == '#,##0," e Ft"'
    assert ws["O5"].number_format == '0" nap"'
    assert ws["R5"].number_format == '#,##0" db"'
    assert ws["B5"].font.b and ws["B5"].font.sz == 14
    assert ws["B6"].value.startswith("MoM ") and "YoY" in ws["B6"].value
    assert ws["B9"].value == "Mutató" and ws["J9"].value == "YTD%"
    assert (
        ws["C10"].number_format == '#,##0," e Ft"' and ws["E10"].number_format == "+0.0%;-0.0%;0.0%"
    )
    assert ws["C12"].number_format == '#,##0" db"'
    assert (
        ws["M9"].value == "Korosítás"
        and ws["M10"].value == "Nem lejárt"
        and ws["M14"].value == "90+ nap"
    )
    assert ws["P10"].number_format == "0.0%" and ws["O10"].number_format == '#,##0" db"'
    assert ws["M15"].value == "Összesen"
    assert ws["M17"].value == "ÁFA-kulcs" and ws["M18"].value == "27%"
    assert ws["B20"].value == "Top-10 vevő" and ws["H20"].value == "Trend (6 hó)"
    assert ws["G21"].number_format == "+0.0%;-0.0%;0.0%"
    exc_header = next(
        r
        for r in range(1, 41)
        if str(ws.cell(row=r, column=13).value or "").startswith("Kivételek")
    )
    assert ws.cell(row=exc_header + 1, column=13).value == "Típus"
    amount_cell = ws.cell(row=exc_header + 2, column=17)
    assert amount_cell.number_format == '#,##0" Ft"'
    assert ws.max_row <= 42
    assert ws["B1"].value == "Havi árbevétel és kintlévőség"
    assert ws["P1"].value == "FIGYELMEZTETÉS (1)"  # planted V07 warn
    assert (
        "2026-08" in ws["B2"].value
        and "spec v1.0.0" in ws["B2"].value
        and "2026-09-03 08:12" in ws["B2"].value
    )
    foot = [ws.cell(row=r, column=2).value for r in range(ws.max_row - 2, ws.max_row + 1)]
    assert (
        foot[0].startswith("Definíciók")
        and "2026-08-31" in foot[1]
        and f"Kit v{KIT_VERSION}" in foot[2]
    )


def test_exceptions_block_capped_and_pointer(built, tmp_path):
    wb = _load(built["out"])
    ws = wb["Vezetői összefoglaló"]
    from report_validate import _exec_exception_rows

    assert _exec_exception_rows(ws) <= rx.MAX_EXEC_EXCEPTIONS
    raw = copy.deepcopy(built["spec"].raw)
    raw["exceptions"] = [
        {
            "id": "all",
            "rule": "overdue_gt_days",
            "params": {"days": 0},
            "severity": "warn",
            "max_rows": 50,
        }
    ]
    spec = rs.spec_from_dict(raw, built["spec"].path)
    result = re_.compute(built["frames"], spec, built["period"])
    assert len(result.exceptions) > rx.MAX_EXEC_EXCEPTIONS
    out = tmp_path / "many.xlsx"
    rx.write_workbook(result, spec, built["period"], out, built["run_info"], frames=built["frames"])
    ws = _load(out)["Vezetői összefoglaló"]
    assert _exec_exception_rows(ws) == rx.MAX_EXEC_EXCEPTIONS
    header = next(
        ws.cell(row=r, column=13).value
        for r in range(1, 41)
        if str(ws.cell(row=r, column=13).value or "").startswith("Kivételek")
    )
    assert header == f"Kivételek (12/{len(result.exceptions)}, teljes lista: Kivételek lap)"
    assert check_output(spec, out).status == "ok"


def test_freeze_panes_merges_and_page_setup(built):
    wb = _load(built["out"])
    ws = wb["Vezetői összefoglaló"]
    assert ws.freeze_panes == "A3"
    assert (
        all(r.min_row == r.max_row == 1 for r in ws.merged_cells.ranges)
        and len(ws.merged_cells.ranges) == 2
    )
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.paperSize == 9
    assert (
        ws.sheet_properties.pageSetUpPr.fitToPage
    )  # fit_to_pages(1, 1); xlsxwriter omits the default 1/1 attrs
    assert ws.column_dimensions["A"].width < 4 and ws.column_dimensions["V"].width < 4
    for name in ("Kivételek", "Számlák", "Ügyfelek", "Mutatók", "Definíciók"):
        assert wb[name].freeze_panes == "A2"
        assert not wb[name].merged_cells.ranges


def test_sparklines_data_bars_and_icons_in_xml(built):
    with zipfile.ZipFile(built["out"]) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    top = len(built["result"].top_customers)
    assert xml.count("<x14:sparkline>") == 6 + top
    assert xml.count("</x14:sparklineGroup>") == 6 + top
    assert xml.count("<dataBar") == 1
    assert xml.count('iconSet="3ArrowsGray"') == 1
    assert "<mergeCell " in xml and xml.count("<mergeCell ") == 2
    assert 'orientation="landscape"' in xml
    assert "#C00000" not in xml  # colours live in styles, red font only via format objects
    with zipfile.ZipFile(built["out"]) as z:
        styles = z.read("xl/styles.xml").decode("utf-8")
    assert "FFC00000" in styles and "FFB45F06" in styles


def test_definitions_and_runlog_sheets(built):
    wb = _load(built["out"])
    spec, period, run_info = built["spec"], built["period"], built["run_info"]
    ws = wb["Definíciók"]
    assert [c.value for c in ws[1]] == [
        "measure_id",
        "label_hu",
        "label_en",
        "formula_words",
        "source_fields",
        "filters_applied",
        "threshold_warn",
        "threshold_critical",
    ]
    assert ws.max_row == len(spec.measures) + 1
    rows = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    assert set(rows) == {m.key for m in spec.measures}
    r = rows["dso"]
    assert (
        ws.cell(row=r, column=2).value.startswith("DSO")
        and ws.cell(row=r, column=4).value == re_.MEASURES["dso"].formula_words
    )
    assert ws.cell(row=r, column=7).value == 45 and ws.cell(row=r, column=8).value == 60
    assert "sztornó a kiállítás hónapjában" in ws.cell(row=r, column=6).value

    ws = wb["Futtatási napló"]
    kv = {
        ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
        for r in range(1, ws.max_row + 1)
    }
    assert kv["Időszak"] == period.label and kv["Futtatva"] == "2026-09-03 08:12"
    assert (
        kv["Spec verzió"] == "1.0.0"
        and len(kv["Spec hash (sha256)"]) == 64
        and kv["Kit verzió"] == KIT_VERSION
    )
    for entity, n in run_info["row_counts"].items():
        assert kv[f"Sorok: {entity}"] == n
    ids = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert [i for i in ids if isinstance(i, str) and i.startswith("V")] == list(rs.CHECK_IDS)
    v07 = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "V07")
    assert ws.cell(row=v07, column=3).value == "figyelmeztetés"
    assert kv["Kézbesített fájl"] == "(kézbesítés még nem történt meg)"


def test_measure_sheet_respects_dimensions(built):
    wb = _load(built["out"])
    ws = wb["Mutatók"]
    headers = [c.value for c in ws[1]]
    dim_col = headers.index("dim_name") + 1
    dims = {ws.cell(row=r, column=dim_col).value for r in range(2, ws.max_row + 1)}
    assert dims <= {"total", "bucket", "month", "customer", "vat_rate"}
    assert "status" in headers


def test_english_and_huf_full_variants(built, tmp_path):
    raw = copy.deepcopy(built["spec"].raw)
    raw["output"]["language"] = "en"
    raw["output"]["number_profile"] = "huf_full"
    raw["output"]["sheets"] = [
        "executive",
        "exceptions",
        "invoice",
        "line",
        "payment",
        "customer",
        "measure",
        "definitions",
        "runlog",
    ]
    spec = rs.spec_from_dict(raw, built["spec"].path)
    out = tmp_path / "en.xlsx"
    rx.write_workbook(
        built["result"], spec, built["period"], out, built["run_info"], frames=built["frames"]
    )
    wb = _load(out)
    assert wb.sheetnames[:-1] == [
        "Executive",
        "Exceptions",
        "Invoices",
        "Lines",
        "Payments",
        "Customers",
        "Measures",
        "Definitions",
        "Run log",
    ]
    ws = wb["Executive"]
    assert (
        ws["B5"].number_format == '#,##0" Ft"'
        and ws["B9"].value == "Measure"
        and ws["M10"].value == "Not due"
    )
    assert ws["P1"].value == "WARNING (1)"
    assert {t for w in wb.worksheets for t in w.tables} == set(rx.expected_tables(spec))
    assert check_output(spec, out).status == "ok"


def test_blocks_order_and_missing_blocks(built, tmp_path):
    raw = copy.deepcopy(built["spec"].raw)
    raw["output"]["executive"]["blocks"] = ["title", "ar_aging", "top_customers", "footnote"]
    spec = rs.spec_from_dict(raw, built["spec"].path)
    out = tmp_path / "blocks.xlsx"
    rx.write_workbook(
        built["result"], spec, built["period"], out, built["run_info"], frames=built["frames"]
    )
    ws = _load(out)["Vezetői összefoglaló"]
    assert ws["B4"].value == "Top-10 vevő"  # no tiles: left column starts at row 4
    assert ws["M4"].value == "Korosítás" and ws["M9"].value == "90+ nap" and ws["M12"].value is None
    with zipfile.ZipFile(out) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert xml.count("<x14:sparkline>") == len(built["result"].top_customers)


def test_empty_result_tables_still_valid(built, tmp_path):
    raw = copy.deepcopy(built["spec"].raw)
    raw["exceptions"] = []
    spec = rs.spec_from_dict(raw, built["spec"].path)
    result = re_.compute(built["frames"], spec, built["period"])
    assert result.exceptions.empty
    out = tmp_path / "empty.xlsx"
    rx.write_workbook(result, spec, built["period"], out, built["run_info"], frames=built["frames"])
    wb = _load(out)
    assert "tbl_havi_arbev_kintlev_exceptions" in wb["Kivételek"].tables
    assert wb["Kivételek"].tables["tbl_havi_arbev_kintlev_exceptions"].ref == "A1:H2"
    assert check_output(spec, out).status == "ok"


# ---------------------------------------------------------------------------
# Power BI CSVs
# ---------------------------------------------------------------------------


def _csv_hashes(built, folder) -> dict[str, str]:
    files = rx.write_powerbi_csvs(built["result"], built["frames"], built["spec"], folder)
    return {p.stem: hashlib.sha256(p.read_bytes()).hexdigest() for p in files}


def print_golden():  # pragma: no cover - maintenance helper
    import tempfile
    from pathlib import Path

    class _F:
        def mktemp(self, name):
            return Path(tempfile.mkdtemp())

    b = built.__wrapped__(_F())
    for k, v in _csv_hashes(b, Path(tempfile.mkdtemp())).items():
        print(f'    "{k}": "{v}",')


def test_powerbi_csvs_bom_headers_and_golden_hashes(built, tmp_path):
    folder = tmp_path / "pbi"
    files = rx.write_powerbi_csvs(built["result"], built["frames"], built["spec"], folder)
    assert [p.name for p in files] == [f"{n}.csv" for n in rs.POWERBI_FILES]
    for p in files:
        data = p.read_bytes()
        assert data.startswith(b"\xef\xbb\xbf")
        header = data.decode("utf-8-sig").splitlines()[0]
        assert ";" not in header and "," in header
    text = {p.stem: p.read_text(encoding="utf-8-sig").splitlines() for p in files}
    assert text["fact_invoice"][0].split(",")[:4] == [
        "invoice_number",
        "chain_root",
        "doc_type",
        "customer_key",
    ]
    assert "as_of_date" in text["fact_invoice"][0] and text["fact_invoice"][1].endswith(
        "2026-08-31"
    )
    assert (
        text["fact_measure"][0]
        == "period,measure_id,dim_name,dim_value,value,prior,prior_year,ytd,ytd_prior,"
        "mom_pct,yoy_pct,avg3m,status"
    )
    assert (
        text["dim_date"][0]
        == "date,period_month,quarter,fiscal_year,fiscal_period_no,is_month_end,is_business_day"
    )
    assert text["dim_date"][1].startswith("2025-07-01,") and text["dim_date"][-1].startswith(
        "2026-08-31,"
    )
    assert (
        text["dim_customer"][0]
        == "customer_key,name,tax_number,country,is_private,first_invoice_date,"
        "last_invoice_date,segment"
    )
    assert (
        text["fact_payment"][0]
        == "payment_id,invoice_number,pay_date,pay_type,amount,currency,amount_huf"
    )
    assert len(text["fact_invoice_line"]) == len(built["frames"].line) + 1
    hashes = {p.stem: hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
    again = _csv_hashes(built, folder)
    assert again == hashes  # overwritten in place, deterministic
    assert hashes == GOLDEN_SHA256


def test_golden_regeneration_helper_signature():
    assert callable(print_golden)
    assert date(2026, 8, 31) == re_.month_end(date(2026, 8, 3))
