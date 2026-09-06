"""import-csv, import-afalista and reconcile: staging tolerance, CSV output, exit codes."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import db
import szamlazz_sync
from conftest import FIXTURES, load_fixture_raw

CSV_HEADER = (
    "Számlaszám;Számla kiállítás dátuma;Vevő neve;Termék neve;Áfakulcs;Tétel nettó érték;"
    "Tétel áfaérték;Tétel bruttó érték;Devizanem"
)


def make_db(tmp_path: Path) -> Path:
    path = tmp_path / "inv.db"
    conn = db.connect(path)
    db.create_schema(conn)
    load_fixture_raw(conn)
    db.rebuild(conn)
    conn.close()
    return path


def run(path: Path, *argv: str) -> int:
    return szamlazz_sync.main(["--db", str(path), *argv])


def read_report(path: Path, period: str) -> list[dict]:
    with (path.parent / f"reconcile-{period}.csv").open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh, delimiter=";"))


def staging(path: Path, table: str) -> list[tuple]:
    conn = db.connect(path)
    try:
        cols = "invoice_number, vat_rate, net, vat, gross"
        return [
            tuple(r)
            for r in conn.execute(
                f"SELECT {cols} FROM {table} ORDER BY invoice_number, vat_rate, id"
            )
        ]
    finally:
        conn.close()


# --- import-csv --------------------------------------------------------------


def test_import_csv_parses_hungarian_formats_and_keeps_extra_columns(tmp_path: Path, capsys):
    path = make_db(tmp_path)
    assert run(path, "import-csv", str(FIXTURES / "fokonyvi_sample.csv")) == 0
    assert "3 sor, 2 számla" in capsys.readouterr().out
    conn = db.connect(path)
    try:
        rows = conn.execute("SELECT * FROM staging_fokonyvi ORDER BY id").fetchall()
        assert len(rows) == 3
        r = rows[0]
        assert (
            r["invoice_number"],
            r["line_no"],
            r["issue_date"],
            r["delivery_date"],
            r["due_date"],
        ) == ("SZLA-2026-1", 1, "2026-08-05", "2026-08-05", "2026-08-20")
        assert (r["customer_name"], r["customer_taxno"], r["product_name"], r["ledger_code"]) == (
            "Alfa Ügyfél Zrt.",
            "23456789-2-41",
            "Szoftverfejlesztés (augusztus)",
            "911",
        )
        assert (r["vat_rate"], r["net"], r["vat"], r["gross"], r["currency"]) == (
            "27",
            100000.0,
            27000.0,
            127000.0,
            "Ft",
        )
        assert (r["pay_status_raw"], r["paid_date"]) == ("100 000,00", "2026-08-18")
        raw = json.loads(r["raw_json"])
        assert raw["Extra oszlop"] == "megjegyzés A"
        assert raw["Kifizetett összeg"] == "100 000,00"
        assert rows[1]["line_no"] == 2 and rows[1]["vat_rate"] == "5"
        assert rows[2]["invoice_number"] == "SZLA-2026-3" and rows[2]["vat"] == 13501.0
        assert (
            conn.execute("SELECT COUNT(*) FROM raw_documents WHERE source = 'csv'").fetchone()[0]
            == 1
        )
        assert (
            conn.execute(
                "SELECT source, fetched FROM sync_log ORDER BY id DESC LIMIT 1"
            ).fetchone()[1]
            == 3
        )
    finally:
        conn.close()
    # re-import replaces rows for the same invoices instead of duplicating them
    assert run(path, "import-csv", str(FIXTURES / "fokonyvi_sample.csv")) == 0
    assert len(staging(path, "staging_fokonyvi")) == 3


def test_import_csv_missing_file_and_missing_column(tmp_path: Path, capsys):
    path = make_db(tmp_path)
    assert run(path, "import-csv", str(tmp_path / "nincs.csv")) == 2
    bad = tmp_path / "bad.csv"
    bad.write_text("Valami;Más\n1;2\n", encoding="utf-8")
    assert run(path, "import-csv", str(bad)) == 2
    assert "számlaszám oszlopot" in capsys.readouterr().out
    tolerant = tmp_path / "tolerant.csv"
    tolerant.write_text(
        'Számlaszám,Tétel nettó érték,Ismeretlen\nSZLA-2026-1,"1 234,5",x\n', encoding="cp1250"
    )
    assert run(path, "import-csv", str(tolerant)) == 0
    rows = staging(path, "staging_fokonyvi")
    assert rows == [("SZLA-2026-1", None, 1234.5, None, None)]


# --- reconcile ---------------------------------------------------------------


def test_reconcile_fixture_csv_one_huf_vat_mismatch(tmp_path: Path, capsys):
    path = make_db(tmp_path)
    assert run(path, "import-csv", str(FIXTURES / "fokonyvi_sample.csv")) == 0
    assert run(path, "reconcile", "--period", "2026-08") == 0
    out = capsys.readouterr().out
    assert "2 számla összevetve" in out and "0 a tűréshatár" in out
    report = read_report(path, "2026-08")
    assert [
        (
            r["invoice_number"],
            r["vat_rate"],
            r["field"],
            r["cache_value"],
            r["staging_value"],
            r["delta"],
        )
        for r in report
    ] == [
        ("SZLA-2026-3", "", "fokonyvi.vat", "13500.00", "13501.00", "1.00"),
        ("SZLA-2026-3", "", "fokonyvi.gross", "63500.00", "63501.00", "1.00"),
        ("SZLA-2026-3", "27", "fokonyvi.vat", "13500.00", "13501.00", "1.00"),
    ]
    assert run(path, "reconcile", "--period", "2026-08", "--tolerance", "0") == 1
    out = capsys.readouterr().out
    assert "3 a tűréshatár (0) felett" in out
    assert "SZLA-2026-3 fokonyvi.vat: cache=13500.00 export=13501.00 delta=+1.00" in out
    assert run(path, "reconcile", "--period", "2026-08", "--tolerance", "0.5") == 1
    assert run(path, "reconcile", "--period", "2026-08", "--tolerance", "1") == 0


def test_reconcile_clean_csv_writes_header_only(tmp_path: Path):
    path = make_db(tmp_path)
    clean = tmp_path / "clean.csv"
    clean.write_text(
        CSV_HEADER + "\n"
        "SZLA-2026-1;2026.08.05.;Alfa Ügyfél Zrt.;Szoftverfejlesztés;27;"
        "100 000,00;27 000,00;127 000,00;Ft\n"
        "SZLA-2026-1;2026.08.05.;Alfa Ügyfél Zrt.;Szakkönyv;5;20 000,00;1 000,00;21 000,00;Ft\n"
        "SZLA-2026-2;2026.08.25.;Alfa Ügyfél Zrt.;Szoftverfejlesztés;27;"
        "-100 000,00;-27 000,00;-127 000,00;Ft\n"
        "SZLA-2026-2;2026.08.25.;Alfa Ügyfél Zrt.;Szakkönyv;5;-20 000,00;-1 000,00;-21 000,00;Ft\n",
        encoding="utf-8",
    )
    assert run(path, "import-csv", str(clean)) == 0
    assert run(path, "reconcile", "--period", "2026-08") == 0
    report = read_report(path, "2026-08")
    assert report == []
    conn = db.connect(path)
    try:
        last = conn.execute(
            "SELECT source, fetched, errors FROM sync_log ORDER BY id DESC LIMIT 1"
        ).fetchone()
        assert tuple(last) == ("reconcile", 2, 0)
    finally:
        conn.close()


def test_reconcile_flags_invoice_missing_from_cache(tmp_path: Path):
    path = make_db(tmp_path)
    extra = tmp_path / "extra.csv"
    extra.write_text(
        CSV_HEADER
        + "\nSZLA-2026-99;2026.08.30.;Ismeretlen Kft.;Valami;27;1 000,00;270,00;1 270,00;Ft\n",
        encoding="utf-8",
    )
    assert run(path, "import-csv", str(extra)) == 0
    assert run(path, "reconcile", "--period", "2026-08") == 1
    report = read_report(path, "2026-08")
    assert len(report) == 1
    assert (
        report[0]["invoice_number"],
        report[0]["field"],
        report[0]["cache_value"],
        report[0]["staging_value"],
        report[0]["delta"],
    ) == ("SZLA-2026-99", "fokonyvi.missing_in_cache", "", "1270.00", "1270.00")


def test_reconcile_without_staging_exits_1(tmp_path: Path, capsys):
    path = make_db(tmp_path)
    assert run(path, "reconcile", "--period", "2026-08") == 1
    assert "import-csv vagy import-afalista" in capsys.readouterr().out


# --- import-afalista ---------------------------------------------------------


def test_import_afalista_unpivots_rates_and_reconciles_clean(tmp_path: Path, capsys):
    path = make_db(tmp_path)
    assert run(path, "import-afalista", str(FIXTURES / "afalista_sample.xlsx")) == 0
    assert "2 számla, 5 sor" in capsys.readouterr().out
    rows = staging(path, "staging_afalista")
    assert rows == [
        ("SZLA-2026-1", "*", 120000.0, 28000.0, 148000.0),
        ("SZLA-2026-1", "27", 100000.0, 27000.0, 127000.0),
        ("SZLA-2026-1", "5", 20000.0, 1000.0, 21000.0),
        ("SZLA-2026-3", "*", 50000.0, 13500.0, 63500.0),
        ("SZLA-2026-3", "27", 50000.0, 13500.0, 63500.0),
    ]
    conn = db.connect(path)
    try:
        r = conn.execute(
            "SELECT issue_date, customer_name, raw_json FROM staging_afalista LIMIT 1"
        ).fetchone()
        assert (r[0], r[1]) == ("2026-08-05", "Alfa Ügyfél Zrt.")
        assert json.loads(r[2])["Fizetési mód"] == "Átutalás"
    finally:
        conn.close()
    assert run(path, "reconcile", "--period", "2026-08", "--tolerance", "0") == 0
    assert read_report(path, "2026-08") == []
    # re-import replaces instead of duplicating
    assert run(path, "import-afalista", str(FIXTURES / "afalista_sample.xlsx")) == 0
    assert len(staging(path, "staging_afalista")) == 5


def test_import_afalista_detects_mismatch_against_cache(tmp_path: Path):
    path = make_db(tmp_path)
    import openpyxl

    wb = openpyxl.load_workbook(FIXTURES / "afalista_sample.xlsx")
    ws = wb.active
    ws["P4"] = 13502  # 27% áfa of SZLA-2026-3
    ws["K4"] = 13502
    ws["L4"] = 63502
    edited = tmp_path / "afalista_edited.xlsx"
    wb.save(edited)
    assert run(path, "import-afalista", str(edited)) == 0
    assert run(path, "reconcile", "--period", "2026-08") == 1
    fields = {
        (r["invoice_number"], r["vat_rate"], r["field"], r["delta"])
        for r in read_report(path, "2026-08")
    }
    assert fields == {
        ("SZLA-2026-3", "", "afalista.vat", "2.00"),
        ("SZLA-2026-3", "", "afalista.gross", "2.00"),
        ("SZLA-2026-3", "27", "afalista.vat", "2.00"),
    }


def test_import_afalista_missing_file(tmp_path: Path):
    assert run(make_db(tmp_path), "import-afalista", str(tmp_path / "nincs.xlsx")) == 2
